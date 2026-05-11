"""
M8: Exportación firmada de resultados.
Genera .xlsx (21 columnas por patente), .json, .csv, .sql INSERT.
Firma SHA-256 del contenido. RF-17..RF-21.
"""
from __future__ import annotations

import csv
import hashlib
import json
from datetime import datetime
from io import BytesIO, StringIO
from pathlib import Path
from typing import Union

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from lingualParam.types import ResultadoCalculo, ResultadoProcesamiento

# 21 columnas según patente LingualParam-CSM™
COLUMNAS_XLSX = [
    "ID_Caso",
    "ID_Paciente_Anonimizado",
    "Fecha_Escaneo",
    "FDI",
    "Prescripcion",
    "T_Actual_deg",
    "A_Actual_deg",
    "I_Actual_deg",
    "T_Objetivo_deg",
    "A_Objetivo_deg",
    "I_Objetivo_deg",
    "Delta_T_deg",
    "Delta_A_deg",
    "Delta_I_deg",
    "Distancia_Cingular_mm",
    "Grosor_Base_mm",
    "Angulo_Slot_deg",
    "Ajuste_Manual",
    "Justificacion_Ajuste",
    "Confianza_Segmentacion",
    "Hash_Dataset",
]

assert len(COLUMNAS_XLSX) == 21, "Deben ser exactamente 21 columnas."


def _calcular_sha256_contenido(contenido: bytes) -> str:
    """RF-21: Calcula SHA-256 del contenido exportado."""
    return hashlib.sha256(contenido).hexdigest()


def _fila_desde_resultado(
    resultado: ResultadoCalculo,
    id_caso: str,
    id_paciente: str,
    fecha_escaneo: str,
    nombre_prescripcion: str,
    hash_dataset: str,
) -> dict:
    """Construye un diccionario con las 21 columnas para un resultado."""
    pieza = resultado.pieza
    return {
        "ID_Caso": id_caso,
        "ID_Paciente_Anonimizado": id_paciente,
        "Fecha_Escaneo": fecha_escaneo,
        "FDI": pieza.fdi,
        "Prescripcion": nombre_prescripcion,
        "T_Actual_deg": pieza.torque_deg,
        "A_Actual_deg": pieza.angulacion_deg,
        "I_Actual_deg": pieza.inclinacion_deg,
        "T_Objetivo_deg": round(pieza.torque_deg + resultado.delta_t, 1),
        "A_Objetivo_deg": round(pieza.angulacion_deg + resultado.delta_a, 1),
        "I_Objetivo_deg": round(pieza.inclinacion_deg + resultado.delta_i, 1),
        "Delta_T_deg": resultado.delta_t,
        "Delta_A_deg": resultado.delta_a,
        "Delta_I_deg": resultado.delta_i,
        "Distancia_Cingular_mm": pieza.distancia_cingular_mm,
        "Grosor_Base_mm": resultado.grosor_base_mm,
        "Angulo_Slot_deg": resultado.angulo_slot_deg,
        "Ajuste_Manual": "Sí" if resultado.ajuste_manual else "No",
        "Justificacion_Ajuste": resultado.justificacion_ajuste,
        "Confianza_Segmentacion": round(pieza.confianza_segmentacion, 3),
        "Hash_Dataset": hash_dataset,
    }


def _construir_filas(
    procesamiento: ResultadoProcesamiento,
    id_paciente: str,
) -> list[dict]:
    fecha_escaneo = datetime.now().strftime("%Y-%m-%d")
    nombre_prescripcion = (
        procesamiento.prescripcion.nombre
        if procesamiento.prescripcion
        else "Sin prescripción"
    )
    return [
        _fila_desde_resultado(
            r,
            procesamiento.id_caso,
            id_paciente,
            fecha_escaneo,
            nombre_prescripcion,
            procesamiento.hash_dataset,
        )
        for r in procesamiento.piezas
    ]


def exportar_xlsx(
    procesamiento: ResultadoProcesamiento,
    id_paciente: str,
    ruta_destino: Union[str, Path],
) -> str:
    """
    RF-18: Exporta el dataset a .xlsx con 21 columnas.
    Retorna el SHA-256 del archivo generado.
    """
    filas = _construir_filas(procesamiento, id_paciente)
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "LingualParam-CSM"

    # Encabezado con formato
    cabecera_relleno = PatternFill(fill_type="solid", fgColor="2563EB")
    cabecera_fuente = Font(bold=True, color="FFFFFF")
    for col_idx, nombre_col in enumerate(COLUMNAS_XLSX, start=1):
        celda = ws.cell(row=1, column=col_idx, value=nombre_col)
        celda.fill = cabecera_relleno
        celda.font = cabecera_fuente
        celda.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = 22

    # Datos
    for row_idx, fila in enumerate(filas, start=2):
        for col_idx, col_nombre in enumerate(COLUMNAS_XLSX, start=1):
            ws.cell(row=row_idx, column=col_idx, value=fila.get(col_nombre, ""))

    # Guardar en buffer para calcular hash
    buffer = BytesIO()
    wb.save(buffer)
    contenido = buffer.getvalue()
    sha256 = _calcular_sha256_contenido(contenido)

    ruta_destino = Path(ruta_destino)
    ruta_destino.write_bytes(contenido)
    return sha256


def exportar_json(
    procesamiento: ResultadoProcesamiento,
    id_paciente: str,
    ruta_destino: Union[str, Path],
) -> str:
    """RF-19: Exporta el dataset a .json para CAD/CAM. Retorna SHA-256."""
    filas = _construir_filas(procesamiento, id_paciente)
    payload = {
        "metadatos": {
            "id_caso": procesamiento.id_caso,
            "id_paciente_anonimizado": id_paciente,
            "prescripcion": (
                procesamiento.prescripcion.nombre if procesamiento.prescripcion else None
            ),
            "fecha_exportacion": datetime.now().isoformat(),
            "version_schema": "1.0",
        },
        "piezas": filas,
    }
    contenido = json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")
    sha256 = _calcular_sha256_contenido(contenido)
    Path(ruta_destino).write_bytes(contenido)
    return sha256


def exportar_csv(
    procesamiento: ResultadoProcesamiento,
    id_paciente: str,
    ruta_destino: Union[str, Path],
) -> str:
    """RF-20: Exporta el dataset a .csv. Retorna SHA-256."""
    filas = _construir_filas(procesamiento, id_paciente)
    buffer = StringIO()
    writer = csv.DictWriter(buffer, fieldnames=COLUMNAS_XLSX, lineterminator="\n")
    writer.writeheader()
    writer.writerows(filas)
    contenido = buffer.getvalue().encode("utf-8-sig")  # BOM para compatibilidad Excel
    sha256 = _calcular_sha256_contenido(contenido)
    Path(ruta_destino).write_bytes(contenido)
    return sha256


def exportar_sql(
    procesamiento: ResultadoProcesamiento,
    id_paciente: str,
    ruta_destino: Union[str, Path],
) -> str:
    """RF-20: Exporta el dataset como sentencias SQL INSERT. Retorna SHA-256."""
    filas = _construir_filas(procesamiento, id_paciente)
    lineas = [
        "-- LingualParam-CSM™ — Exportación SQL",
        f"-- Caso: {procesamiento.id_caso}",
        f"-- Fecha: {datetime.now().isoformat()}",
        "",
    ]
    for fila in filas:
        cols = ", ".join(COLUMNAS_XLSX)
        def escapar(v: object) -> str:
            if v is None:
                return "NULL"
            s = str(v).replace("'", "''")
            return f"'{s}'"
        vals = ", ".join(escapar(fila.get(c)) for c in COLUMNAS_XLSX)
        lineas.append(f"INSERT INTO Pieza ({cols}) VALUES ({vals});")

    contenido = "\n".join(lineas).encode("utf-8")
    sha256 = _calcular_sha256_contenido(contenido)
    Path(ruta_destino).write_bytes(contenido)
    return sha256


def exportar_todo(
    procesamiento: ResultadoProcesamiento,
    id_paciente: str,
    directorio_destino: Union[str, Path],
) -> dict[str, str]:
    """
    RF-17..RF-21: Exporta todos los formatos y retorna dict {formato: sha256}.
    El directorio debe existir previamente (confirmación explícita del usuario RF-17).
    """
    dir_dst = Path(directorio_destino)
    id_caso = procesamiento.id_caso
    hashes: dict[str, str] = {}

    hashes["xlsx"] = exportar_xlsx(procesamiento, id_paciente, dir_dst / f"{id_caso}.xlsx")
    hashes["json"] = exportar_json(procesamiento, id_paciente, dir_dst / f"{id_caso}.json")
    hashes["csv"] = exportar_csv(procesamiento, id_paciente, dir_dst / f"{id_caso}.csv")
    hashes["sql"] = exportar_sql(procesamiento, id_paciente, dir_dst / f"{id_caso}.sql")

    # Hash global del dataset (sobre concatenación de hashes individuales)
    hashes["dataset"] = _calcular_sha256_contenido(
        "".join(hashes.values()).encode("utf-8")
    )

    return hashes
