"""
Tests para M8: Exportación firmada (xlsx, json, csv, sql) + SHA-256.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

import openpyxl
import pytest
import numpy as np

from lingualParam.m8_exporter import (
    COLUMNAS_XLSX,
    _calcular_sha256_contenido,
    exportar_csv,
    exportar_json,
    exportar_sql,
    exportar_todo,
    exportar_xlsx,
)
from lingualParam.types import (
    EstadoCaso,
    MallaImportada,
    PiezaDental,
    PrescripcionCompleta,
    ResultadoCalculo,
    ResultadoProcesamiento,
    TipoPrescripcion,
)


def _resultado_sintetico(fdi: int = 11) -> ResultadoCalculo:
    """Crea un ResultadoCalculo mínimo para pruebas."""
    vertices = np.zeros((100, 3))
    pieza = PiezaDental(
        fdi=fdi,
        vertices=vertices,
        centroide=np.zeros(3),
        centroide_lingual=np.array([0.0, -3.0, 0.0]),
        eje_longitudinal=np.array([0.0, 0.0, 1.0]),
        vector_normal_lingual=np.array([0.0, -1.0, 0.0]),
        torque_deg=5.5,
        angulacion_deg=3.0,
        inclinacion_deg=-2.0,
        distancia_cingular_mm=0.8,
        confianza_segmentacion=0.97,
    )
    return ResultadoCalculo(
        pieza=pieza,
        prescripcion_fdi=fdi,
        delta_t=1.5,
        delta_a=-0.5,
        delta_i=2.0,
        grosor_base_mm=0.6,
        angulo_slot_deg=1.5,
        ajuste_manual=False,
        justificacion_ajuste="",
    )


def _procesamiento_sintetico() -> ResultadoProcesamiento:
    """Crea un ResultadoProcesamiento mínimo."""
    malla = MallaImportada(
        ruta_origen="/test/modelo.stl",
        formato="stl",
        vertices=np.zeros((100, 3)),
        caras=np.zeros((50, 3), dtype=np.int32),
        unidad_mm=True,
        arcada="superior",
        hash_origen="abc123",
    )
    prescripcion = PrescripcionCompleta(
        nombre="STb",
        tipo=TipoPrescripcion.STB,
        descripcion="Prescripción STb Scuzzo-Takemoto",
        referencia_bibliografica="Scuzzo & Takemoto, 2003",
    )
    return ResultadoProcesamiento(
        id_caso="caso-test-001",
        estado=EstadoCaso.PROCESADO,
        malla=malla,
        piezas=[_resultado_sintetico(11), _resultado_sintetico(21)],
        prescripcion=prescripcion,
        hash_dataset="",
    )


class TestSHA256:
    def test_sha256_es_64_chars(self) -> None:
        h = _calcular_sha256_contenido(b"test")
        assert len(h) == 64

    def test_sha256_deterministico(self) -> None:
        contenido = b"contenido de prueba"
        assert _calcular_sha256_contenido(contenido) == _calcular_sha256_contenido(contenido)

    def test_sha256_distinto_para_distintos_contenidos(self) -> None:
        assert _calcular_sha256_contenido(b"A") != _calcular_sha256_contenido(b"B")


class TestExportarXLSX:
    def test_crea_archivo_xlsx(self, tmp_path: Path) -> None:
        proc = _procesamiento_sintetico()
        ruta = tmp_path / "resultado.xlsx"
        exportar_xlsx(proc, "id_paciente_test", str(ruta))
        assert ruta.exists()

    def test_xlsx_tiene_21_columnas(self, tmp_path: Path) -> None:
        proc = _procesamiento_sintetico()
        ruta = tmp_path / "resultado.xlsx"
        exportar_xlsx(proc, "id_paciente_test", str(ruta))
        wb = openpyxl.load_workbook(str(ruta))
        ws = wb.active
        encabezados = [ws.cell(row=1, column=i).value for i in range(1, 22)]
        assert encabezados == COLUMNAS_XLSX

    def test_xlsx_retorna_sha256_valido(self, tmp_path: Path) -> None:
        proc = _procesamiento_sintetico()
        ruta = tmp_path / "resultado.xlsx"
        sha = exportar_xlsx(proc, "id_paciente_test", str(ruta))
        assert len(sha) == 64

    def test_xlsx_tiene_filas_por_pieza(self, tmp_path: Path) -> None:
        proc = _procesamiento_sintetico()
        ruta = tmp_path / "resultado.xlsx"
        exportar_xlsx(proc, "id_paciente_test", str(ruta))
        wb = openpyxl.load_workbook(str(ruta))
        ws = wb.active
        # 1 fila encabezado + 2 piezas
        assert ws.max_row == 3


class TestExportarJSON:
    def test_crea_archivo_json(self, tmp_path: Path) -> None:
        proc = _procesamiento_sintetico()
        ruta = tmp_path / "resultado.json"
        exportar_json(proc, "id_paciente_test", str(ruta))
        assert ruta.exists()

    def test_json_tiene_estructura_correcta(self, tmp_path: Path) -> None:
        proc = _procesamiento_sintetico()
        ruta = tmp_path / "resultado.json"
        exportar_json(proc, "id_paciente_test", str(ruta))
        datos = json.loads(ruta.read_text(encoding="utf-8"))
        assert "metadatos" in datos
        assert "piezas" in datos
        assert len(datos["piezas"]) == 2


class TestExportarCSV:
    def test_crea_archivo_csv(self, tmp_path: Path) -> None:
        proc = _procesamiento_sintetico()
        ruta = tmp_path / "resultado.csv"
        exportar_csv(proc, "id_paciente_test", str(ruta))
        assert ruta.exists()

    def test_csv_tiene_21_columnas_en_encabezado(self, tmp_path: Path) -> None:
        proc = _procesamiento_sintetico()
        ruta = tmp_path / "resultado.csv"
        exportar_csv(proc, "id_paciente_test", str(ruta))
        with open(str(ruta), encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            assert list(reader.fieldnames) == COLUMNAS_XLSX


class TestExportarSQL:
    def test_crea_archivo_sql(self, tmp_path: Path) -> None:
        proc = _procesamiento_sintetico()
        ruta = tmp_path / "resultado.sql"
        exportar_sql(proc, "id_paciente_test", str(ruta))
        assert ruta.exists()

    def test_sql_contiene_insert_statements(self, tmp_path: Path) -> None:
        proc = _procesamiento_sintetico()
        ruta = tmp_path / "resultado.sql"
        exportar_sql(proc, "id_paciente_test", str(ruta))
        contenido = ruta.read_text(encoding="utf-8")
        assert "INSERT INTO Pieza" in contenido

    def test_sql_tiene_una_linea_por_pieza(self, tmp_path: Path) -> None:
        proc = _procesamiento_sintetico()
        ruta = tmp_path / "resultado.sql"
        exportar_sql(proc, "id_paciente_test", str(ruta))
        contenido = ruta.read_text(encoding="utf-8")
        n_inserts = contenido.count("INSERT INTO")
        assert n_inserts == len(proc.piezas)


class TestExportarTodo:
    def test_crea_todos_los_formatos(self, tmp_path: Path) -> None:
        proc = _procesamiento_sintetico()
        hashes = exportar_todo(proc, "id_paciente_test", str(tmp_path))
        assert (tmp_path / "caso-test-001.xlsx").exists()
        assert (tmp_path / "caso-test-001.json").exists()
        assert (tmp_path / "caso-test-001.csv").exists()
        assert (tmp_path / "caso-test-001.sql").exists()

    def test_retorna_dict_con_hashes(self, tmp_path: Path) -> None:
        proc = _procesamiento_sintetico()
        hashes = exportar_todo(proc, "id_paciente_test", str(tmp_path))
        assert set(hashes.keys()) == {"xlsx", "json", "csv", "sql", "dataset"}
        for sha in hashes.values():
            assert len(sha) == 64
